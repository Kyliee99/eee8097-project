"""JSON, CSV and XLSX result export."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any


SUMMARY_FIELDS = [
    "model", "operation", "accuracy", "macro_precision", "macro_recall", "macro_f1",
    "weighted_f1", "acev2_theoretical", "acev2_actual",
    "inference_seconds_mean", "seconds_per_image", "throughput_images_per_second",
    "checkpoint_size_mb", "deployment_size_mb", "packed_xnor_word_ops",
    "tm_logic_ops_upper_bound", "tm_clause_length_mean", "tm_patches_per_image",
    "tm_all_patches_visited", "estimated_energy_pj", "energy_status", "cost_status",
]


def _acev2_values(result: dict[str, Any]) -> tuple[Any, Any]:
    """Return comparable theoretical and actual ACEv2 values for one model."""
    cost = result.get("cost", {})
    if result.get("model") == "tsetlin":
        return (
            cost.get("extended_acev2_tm_upper_bound"),
            cost.get("extended_acev2_tm_trained_model"),
        )
    algorithmic = cost.get("algorithmic_acev2") or {}
    python_tensor = cost.get("python_tensor_acev2") or {}
    return algorithmic.get("acev2"), python_tensor.get("acev2")


def _summary_row(result: dict[str, Any]) -> dict[str, Any]:
    cost = result.get("cost", {})
    clause_profile = cost.get("trained_clause_profile") or {}
    patch_evaluation = cost.get("patch_evaluation") or {}
    acev2_theoretical, acev2_actual = _acev2_values(result)
    return {
        "model": result.get("model"),
        "operation": result.get("operation"),
        "accuracy": result.get("metrics", {}).get("accuracy"),
        "macro_precision": result.get("metrics", {}).get("macro_precision"),
        "macro_recall": result.get("metrics", {}).get("macro_recall"),
        "macro_f1": result.get("metrics", {}).get("macro_f1"),
        "weighted_f1": result.get("metrics", {}).get("weighted_f1"),
        "acev2_theoretical": acev2_theoretical,
        "acev2_actual": acev2_actual,
        "inference_seconds_mean": result.get("timing", {}).get("inference_seconds_mean"),
        "seconds_per_image": result.get("timing", {}).get("seconds_per_image"),
        "throughput_images_per_second": result.get("timing", {}).get("throughput_images_per_second"),
        "checkpoint_size_mb": result.get("storage", {}).get("checkpoint_size_mb"),
        "deployment_size_mb": result.get("storage", {}).get("deployment_size_mb"),
        "packed_xnor_word_ops": cost.get("packed_xnor_word_operations"),
        "tm_logic_ops_upper_bound": cost.get("logic_operations_upper_bound"),
        "tm_clause_length_mean": clause_profile.get("included_literals_mean"),
        "tm_patches_per_image": cost.get("patches_per_image"),
        "tm_all_patches_visited": patch_evaluation.get("all_patches_visited"),
        "estimated_energy_pj": result.get("energy", {}).get("estimated_energy_pj"),
        "energy_status": result.get("energy", {}).get("energy_status"),
        "cost_status": cost.get("acev2_status"),
    }


def export_results(payload: dict[str, Any], output_dir: str | Path) -> dict[str, Path]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    paths = {
        "json": output / "comparison_results.json",
        "csv": output / "comparison_summary.csv",
        "xlsx": output / "comparison_results.xlsx",
    }
    with paths["json"].open("w", encoding="utf-8") as stream:
        json.dump(payload, stream, indent=2, ensure_ascii=False)
    rows = [_summary_row(result) for result in payload["results"]]
    with paths["csv"].open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=SUMMARY_FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        paths.pop("xlsx")
        return paths
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(SUMMARY_FIELDS)
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        summary.append([row[field] for field in SUMMARY_FIELDS])
    for column in summary.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
        summary.column_dimensions[column[0].column_letter].width = width
    workbook.save(paths["xlsx"])
    return paths

