# -*- coding: utf-8 -*-
"""
公共模型比较文件。

默认比较四种算法场景：
1. MLP
2. BNN（BinaryLinear 二值点积）
3. BNN（XNOR + Popcount 推理）
4. Tsetlin Machine

所有模型使用同一个 RGBDObjectDataset 测试集。
默认使用 CPU，便于与 CPU Tsetlin Machine 进行较公平的比较。

运行：
    python compare.py

只比较指定模型：
    python compare.py --models mlp bnn
    python compare.py --models mlp bnn bnn_xnor tsetlin

指定模型路径：
    python compare.py ^
        --mlp-model models/rgbd_mlp_2.pth ^
        --bnn-model models/rgbd_bnn_2.pth ^
        --tsetlin-model tsetlin_output/best_tsetlin_rgbd.pkl
"""

from __future__ import annotations

import argparse
import csv
import json
import pickle
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import torch
import torch.nn as nn
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.metrics import classification_report

from bnn import RGBDBNN
from common import (
    BATCH_SIZE,
    DATASET_ROOT,
    IMAGE_SIZE,
    build_classification_metrics,
    create_dataloaders,
    set_seed,
)
from mlp import RGBDMLP


SCRIPT_DIR = Path(__file__).resolve().parent

DEFAULT_MLP_MODEL = (
    SCRIPT_DIR
    / "models"
    / "rgbd_mlp_2.pth"
)

DEFAULT_BNN_MODEL = (
    SCRIPT_DIR
    / "models"
    / "rgbd_bnn_2.pth"
)

DEFAULT_TSETLIN_MODEL = (
    SCRIPT_DIR
    / "tsetlin5_output"
    / "best_tsetlin_rgbd_rgb5_2.pkl"
)

DEFAULT_OUTPUT_DIR = (
    SCRIPT_DIR
    / "comparison_results"
)

SUPPORTED_MODELS = (
    "mlp",
    "bnn",
    "bnn_xnor",
    "tsetlin",
)


# ============================================================
# 1. 命令行参数
# ============================================================

def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Compare MLP, BNN, BNN-XNOR/Popcount "
            "and Tsetlin Machine on the same test set."
        )
    )

    parser.add_argument(
        "--models",
        nargs="+",
        choices=SUPPORTED_MODELS,
        default=list(SUPPORTED_MODELS),
        help=(
            "要比较的模型。默认："
            "mlp bnn bnn_xnor tsetlin"
        ),
    )

    parser.add_argument(
        "--dataset-root",
        type=Path,
        default=DATASET_ROOT,
        help="RGB-D Object Dataset 根目录。",
    )

    parser.add_argument(
        "--batch-size",
        type=int,
        default=BATCH_SIZE,
        help="测试 batch size。",
    )

    parser.add_argument(
        "--device",
        choices=("cpu", "cuda", "auto"),
        default="cpu",
        help=(
            "PyTorch模型使用的设备。默认CPU，"
            "便于与CPU Tsetlin公平比较。"
        ),
    )

    parser.add_argument(
        "--mlp-model",
        type=Path,
        default=DEFAULT_MLP_MODEL,
        help="MLP .pth 文件。",
    )

    parser.add_argument(
        "--bnn-model",
        type=Path,
        default=DEFAULT_BNN_MODEL,
        help="BNN .pth 文件。",
    )

    parser.add_argument(
        "--tsetlin-model",
        type=Path,
        default=DEFAULT_TSETLIN_MODEL,
        help="Tsetlin .pkl 文件。",
    )

    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="比较结果保存目录。",
    )

    return parser.parse_args()


def resolve_device(device_name: str) -> torch.device:
    if device_name == "cpu":
        return torch.device("cpu")

    if device_name == "cuda":
        if not torch.cuda.is_available():
            raise RuntimeError(
                "指定了CUDA，但当前环境没有可用CUDA设备。"
            )
        return torch.device("cuda")

    return torch.device(
        "cuda"
        if torch.cuda.is_available()
        else "cpu"
    )


# ============================================================
# 2. 模型参数加载
# ============================================================

def load_torch_weights(
    model: nn.Module,
    model_path: Path,
    device: torch.device,
) -> None:
    """
    兼容三种常见保存形式：

    1. torch.save(model.state_dict(), path)
    2. {"model_state_dict": state_dict}
    3. {"state_dict": state_dict}
    """
    if not model_path.exists():
        raise FileNotFoundError(
            f"模型文件不存在：{model_path.resolve()}"
        )

    checkpoint = torch.load(
        model_path,
        map_location=device,
        weights_only=False,
    )

    if isinstance(checkpoint, nn.Module):
        state_dict = checkpoint.state_dict()

    elif isinstance(checkpoint, dict):
        if "model_state_dict" in checkpoint:
            state_dict = checkpoint[
                "model_state_dict"
            ]
        elif "state_dict" in checkpoint:
            state_dict = checkpoint["state_dict"]
        else:
            state_dict = checkpoint

    else:
        raise ValueError(
            f"无法识别模型文件格式：{model_path}"
        )

    model.load_state_dict(state_dict)
    model.eval()


# ============================================================
# 3. 测试数据缓存
# ============================================================

def cache_test_batches(
    testloader,
) -> tuple[
    list[tuple[torch.Tensor, torch.Tensor]],
    np.ndarray,
]:
    """
    先把测试集读取到内存。

    这样计时时不包含硬盘读取和PIL预处理时间，
    更接近模型自身的推理时间。
    """
    cached_batches = []
    labels_all = []

    for images, labels in testloader:
        cached_batches.append((
            images.contiguous(),
            labels.contiguous(),
        ))
        labels_all.append(
            labels.numpy()
        )

    if not cached_batches:
        raise RuntimeError(
            "测试集中没有读取到任何数据。"
        )

    return (
        cached_batches,
        np.concatenate(
            labels_all,
            axis=0,
        ),
    )


# ============================================================
# 4. PyTorch模型测试
# ============================================================

def evaluate_torch_model(
    model: nn.Module,
    cached_batches,
    class_names: list[str],
    device: torch.device,
) -> dict[str, Any]:
    predictions_all = []
    labels_all = []

    model.eval()

    # 预热一次，减少首次调用初始化对时间的影响。
    with torch.no_grad():
        first_images, _ = cached_batches[0]
        _ = model(
            first_images.to(device)
        )

    if device.type == "cuda":
        torch.cuda.synchronize()

    start_time = time.perf_counter()

    with torch.no_grad():
        for images, labels in cached_batches:
            images = images.to(
                device,
                non_blocking=True,
            )

            outputs = model(images)

            predictions = torch.argmax(
                outputs,
                dim=1,
            )

            predictions_all.extend(
                predictions.cpu().numpy()
            )
            labels_all.extend(
                labels.numpy()
            )

    if device.type == "cuda":
        torch.cuda.synchronize()

    total_time = (
        time.perf_counter()
        - start_time
    )

    return build_classification_metrics(
        labels_all,
        predictions_all,
        class_names,
        total_time=total_time,
    )


# ============================================================
# 5. Tsetlin测试数据转换
# ============================================================

def load_tsetlin_artifact(
    model_path: Path,
    class_names: list[str],
) -> dict[str, Any]:
    """
    加载Tsetlin模型及其预处理元数据。

    支持：
    1. TM5.py保存的checkpoint字典；
    2. 常见的model/tsetlin_model/classifier键；
    3. 直接pickle.dump(model)的旧格式。

    比较文件不读取或硬编码Tsetlin内部模型shape。
    输入Boolean通道数量由checkpoint中的阈值结构动态生成。
    """
    if not model_path.exists():
        raise FileNotFoundError(
            "Tsetlin模型文件不存在："
            f"{model_path.resolve()}"
        )

    # pickle只能加载可信的模型文件。
    with model_path.open("rb") as file:
        checkpoint = pickle.load(file)

    checkpoint_data: dict[str, Any]

    if isinstance(checkpoint, dict):
        checkpoint_data = checkpoint

        model = None

        for model_key in (
            "model",
            "tsetlin_model",
            "classifier",
        ):
            candidate = checkpoint.get(model_key)

            if (
                candidate is not None
                and hasattr(candidate, "predict")
            ):
                model = candidate
                break

        if model is None:
            raise KeyError(
                "Tsetlin checkpoint中没有可识别的模型对象。"
                "需要包含'model'、'tsetlin_model'或"
                "'classifier'，并且对象具有predict()方法。"
            )

    elif hasattr(checkpoint, "predict"):
        # 兼容直接pickle.dump(model)
        model = checkpoint
        checkpoint_data = {}

    else:
        raise TypeError(
            "无法识别Tsetlin模型文件格式："
            f"{type(checkpoint).__name__}"
        )

    saved_classes = checkpoint_data.get(
        "class_names"
    )

    if (
        saved_classes is not None
        and list(saved_classes) != list(class_names)
    ):
        raise RuntimeError(
            "Tsetlin模型类别顺序与测试集不一致。\n"
            f"模型类别：{list(saved_classes)}\n"
            f"测试类别：{list(class_names)}"
        )

    threshold_source = None
    raw_thresholds = None

    # 优先读取TM5.py使用的字段。
    for threshold_key in (
        "rgb_channel_thresholds",
        "channel_thresholds",
        "boolean_thresholds",
        "thresholds",
    ):
        if checkpoint_data.get(threshold_key) is not None:
            threshold_source = threshold_key
            raw_thresholds = checkpoint_data[
                threshold_key
            ]
            break

    input_mode = checkpoint_data.get(
        "input_mode"
    )

    if input_mode not in {
        "rgb",
        "grayscale",
    }:
        # 旧checkpoint没有input_mode时，只根据阈值外层结构推断。
        if (
            isinstance(raw_thresholds, (list, tuple))
            and len(raw_thresholds) == 1
            and isinstance(
                raw_thresholds[0],
                (list, tuple, np.ndarray),
            )
        ):
            input_mode = "grayscale"
        else:
            input_mode = "rgb"

    metadata = {
        "best_epoch": checkpoint_data.get(
            "epoch"
        ),
        "best_validation_accuracy": (
            checkpoint_data.get(
                "validation_accuracy"
            )
        ),
        "best_validation_macro_f1": (
            checkpoint_data.get(
                "validation_macro_f1"
            )
        ),
        "patch_dim": checkpoint_data.get(
            "patch_dim"
        ),
        "number_of_clauses": checkpoint_data.get(
            "number_of_clauses"
        ),
        "T": checkpoint_data.get("T"),
        "s": checkpoint_data.get("s"),
        "max_included_literals": (
            checkpoint_data.get(
                "max_included_literals"
            )
        ),
        "weighted_clauses": checkpoint_data.get(
            "weighted_clauses"
        ),
        "input_mode": input_mode,
        "threshold_source": threshold_source,
        # 仅保存checkpoint中的原始记录，用于诊断。
        # 真正的Boolean通道数会在转换数据后动态计算。
        "checkpoint_boolean_channels": (
            checkpoint_data.get(
                "boolean_channels"
            )
        ),
    }

    return {
        "model": model,
        "input_mode": input_mode,
        "raw_thresholds": raw_thresholds,
        "metadata": metadata,
    }


def _normalise_channel_thresholds(
    raw_thresholds,
    channel_count: int,
) -> tuple[tuple[float, ...], ...]:
    """
    将不同checkpoint格式统一成：

        (
            (channel_0_threshold_0, ...),
            (channel_1_threshold_0, ...),
            ...
        )

    支持：
    - TM5格式：3个RGB通道，每个通道任意数量阈值；
    - 每个通道只有一个阈值的扁平列表；
    - 所有通道共享同一组阈值的扁平列表；
    - 旧模型没有阈值元数据时，每个通道默认使用0.5。
    """
    if channel_count <= 0:
        raise ValueError(
            f"无效输入通道数量：{channel_count}"
        )

    if raw_thresholds is None:
        print(
            "Warning: Tsetlin checkpoint没有保存阈值。"
            "将为每个输入通道使用默认阈值0.5。"
        )

        return tuple(
            (0.5,)
            for _ in range(channel_count)
        )

    if isinstance(raw_thresholds, np.ndarray):
        raw_thresholds = raw_thresholds.tolist()

    if not isinstance(
        raw_thresholds,
        (list, tuple),
    ):
        raise TypeError(
            "Tsetlin阈值必须是list、tuple或NumPy数组，"
            f"当前为{type(raw_thresholds).__name__}。"
        )

    if len(raw_thresholds) == 0:
        raise ValueError(
            "Tsetlin阈值列表不能为空。"
        )

    is_nested = all(
        isinstance(
            value,
            (list, tuple, np.ndarray),
        )
        for value in raw_thresholds
    )

    if is_nested:
        if len(raw_thresholds) != channel_count:
            raise ValueError(
                "checkpoint中的阈值通道数与测试图片不一致："
                f"阈值包含{len(raw_thresholds)}个通道，"
                f"图片包含{channel_count}个通道。"
            )

        channel_thresholds = []

        for channel_index, values in enumerate(
            raw_thresholds
        ):
            if isinstance(values, np.ndarray):
                values = values.tolist()

            if len(values) == 0:
                raise ValueError(
                    f"第{channel_index}个通道没有阈值。"
                )

            channel_thresholds.append(
                tuple(
                    float(value)
                    for value in values
                )
            )

        return tuple(channel_thresholds)

    flat_thresholds = tuple(
        float(value)
        for value in raw_thresholds
    )

    if len(flat_thresholds) == channel_count:
        # 兼容旧格式：(R_threshold, G_threshold, B_threshold)
        return tuple(
            (threshold,)
            for threshold in flat_thresholds
        )

    # 其他扁平格式解释为：所有输入通道共享同一组阈值。
    return tuple(
        flat_thresholds
        for _ in range(channel_count)
    )


def tensor_to_boolean_dynamic(
    images: torch.Tensor,
    channel_thresholds: tuple[
        tuple[float, ...],
        ...
    ],
) -> np.ndarray:
    """
    根据阈值元数据动态生成Tsetlin Boolean输入。

    输入：
        images：[B,C,H,W]

    输出：
        [B,H,W,total_boolean_channels]

    total_boolean_channels由每个输入通道的阈值数量之和决定，
    不硬编码为3、9、21或其他固定值。
    """
    if images.ndim != 4:
        raise ValueError(
            "Tsetlin图片输入应为[B,C,H,W]，"
            f"当前为{tuple(images.shape)}"
        )

    channel_count = int(images.shape[1])

    if len(channel_thresholds) != channel_count:
        raise ValueError(
            "阈值通道数与图片通道数不一致："
            f"{len(channel_thresholds)} != {channel_count}"
        )

    boolean_planes = []

    for channel_index, thresholds in enumerate(
        channel_thresholds
    ):
        threshold_tensor = torch.tensor(
            thresholds,
            dtype=images.dtype,
            device=images.device,
        ).view(
            1,
            len(thresholds),
            1,
            1,
        )

        # [B,H,W] -> [B,1,H,W]
        channel_images = images[
            :,
            channel_index,
            :,
            :,
        ].unsqueeze(1)

        # [B,number_of_thresholds,H,W]
        channel_boolean_planes = (
            channel_images > threshold_tensor
        )

        boolean_planes.append(
            channel_boolean_planes
        )

    # [B,total_boolean_channels,H,W]
    boolean_tensor = torch.cat(
        boolean_planes,
        dim=1,
    ).to(torch.uint32)

    # TMU卷积输入：[B,H,W,Boolean通道]
    boolean_tensor = boolean_tensor.permute(
        0,
        2,
        3,
        1,
    ).contiguous()

    return (
        boolean_tensor
        .cpu()
        .numpy()
        .astype(
            np.uint32,
            copy=False,
        )
    )


def cached_batches_to_boolean_numpy(
    cached_batches,
    raw_thresholds,
) -> tuple[
    np.ndarray,
    np.ndarray,
    dict[str, Any],
]:
    """
    将缓存的PyTorch batch转换为Tsetlin NumPy输入。

    不查询Tsetlin模型内部shape，而是根据：
    1. 实际测试图片通道数；
    2. checkpoint保存的阈值结构；
    动态构造Boolean输入。
    """
    if not cached_batches:
        raise RuntimeError(
            "测试集中没有缓存任何batch。"
        )

    first_images, _ = cached_batches[0]

    if first_images.ndim != 4:
        raise ValueError(
            "测试图片应为[B,C,H,W]，"
            f"当前为{tuple(first_images.shape)}"
        )

    channel_count = int(
        first_images.shape[1]
    )

    channel_thresholds = (
        _normalise_channel_thresholds(
            raw_thresholds,
            channel_count,
        )
    )

    feature_batches = []
    label_batches = []

    with torch.no_grad():
        for images, labels in cached_batches:

            if images.ndim != 4:
                raise ValueError(
                    "测试图片应为[B,C,H,W]，"
                    f"当前为{tuple(images.shape)}"
                )

            if int(images.shape[1]) != channel_count:
                raise ValueError(
                    "不同batch的图片通道数量不一致。"
                )

            feature_batches.append(
                tensor_to_boolean_dynamic(
                    images,
                    channel_thresholds,
                )
            )

            label_batches.append(
                labels
                .cpu()
                .numpy()
                .astype(
                    np.uint32,
                    copy=False,
                )
            )

    X_test = np.concatenate(
        feature_batches,
        axis=0,
    )

    y_test = np.concatenate(
        label_batches,
        axis=0,
    )

    boolean_channels = sum(
        len(thresholds)
        for thresholds in channel_thresholds
    )

    preprocessing_metadata = {
        "channel_thresholds": [
            list(thresholds)
            for thresholds in channel_thresholds
        ],
        # 保留旧Excel字段名的兼容性。
        "boolean_thresholds": [
            list(thresholds)
            for thresholds in channel_thresholds
        ],
        "input_channels": channel_count,
        "boolean_channels": boolean_channels,
        "generated_tsetlin_input_shape": list(
            X_test.shape[1:]
        ),
    }

    return (
        np.ascontiguousarray(
            X_test,
            dtype=np.uint32,
        ),
        np.ascontiguousarray(
            y_test,
            dtype=np.uint32,
        ),
        preprocessing_metadata,
    )


def evaluate_tsetlin_model(
    artifact: dict[str, Any],
    cached_batches,
    class_names: list[str],
) -> tuple[dict[str, Any], dict[str, Any]]:
    """
    测试已经加载的Tsetlin模型。

    该函数不读取模型内部shape，也不写死Boolean通道数。
    模型输入由checkpoint预处理元数据动态生成。
    """
    model = artifact["model"]

    (
        X_test,
        y_test,
        preprocessing_metadata,
    ) = cached_batches_to_boolean_numpy(
        cached_batches,
        artifact.get("raw_thresholds"),
    )

    print(
        "Generated Tsetlin test shape:",
        X_test.shape,
    )
    print(
        "Generated Boolean channels:",
        preprocessing_metadata[
            "boolean_channels"
        ],
    )

    # 预热一次。
    warmup_count = min(
        1,
        len(X_test),
    )

    try:
        if warmup_count > 0:
            _ = model.predict(
                X_test[:warmup_count]
            )

    except Exception as error:
        raise RuntimeError(
            "Tsetlin模型无法接收动态生成的测试输入。\n"
            f"生成的输入shape：{X_test.shape}\n"
            "请确认该.pkl确实由当前TM5预处理方式训练，"
            "并且checkpoint保存了正确的阈值。"
        ) from error

    start_time = time.perf_counter()

    predictions = np.asarray(
        model.predict(X_test)
    ).reshape(-1)

    total_time = (
        time.perf_counter()
        - start_time
    )

    if len(predictions) != len(y_test):
        raise RuntimeError(
            "Tsetlin预测数量与测试标签数量不一致："
            f"{len(predictions)} != {len(y_test)}"
        )

    metrics = build_classification_metrics(
        y_test,
        predictions,
        class_names,
        total_time=total_time,
    )

    model_metadata = {
        **artifact.get("metadata", {}),
        **preprocessing_metadata,
    }

    return metrics, model_metadata


# ============================================================
# 6. 参数、权重和文件大小
# ============================================================

def count_total_parameters(
    model: nn.Module,
) -> int:
    return sum(
        parameter.numel()
        for parameter in model.parameters()
    )


def count_matrix_weights(
    model: nn.Module,
) -> int:
    """
    只统计Linear等二维权重，
    不统计bias和BatchNorm的一维参数。
    """
    return sum(
        parameter.numel()
        for name, parameter
        in model.named_parameters()
        if (
            "weight" in name
            and parameter.dim() >= 2
        )
    )


def get_file_size_mb(
    file_path: Path,
) -> float:
    if not file_path.exists():
        return 0.0

    return (
        file_path.stat().st_size
        / (1024 * 1024)
    )


def theoretical_size_mb(
    weight_count: int,
    bits_per_weight: int,
) -> float:
    return (
        weight_count
        * bits_per_weight
        / 8
        / (1024 * 1024)
    )


def add_common_statistics(
    model_name: str,
    model_path: Path,
    metrics: dict[str, Any],
    model: nn.Module | None = None,
    bits_per_weight: int | None = None,
    operation: str = "",
) -> dict[str, Any]:
    result = {
        "model": model_name,
        "operation": operation,
        **metrics,
        "model_file_size_mb": (
            get_file_size_mb(
                model_path
            )
        ),
    }

    if model is not None:
        total_parameters = (
            count_total_parameters(model)
        )
        matrix_weights = (
            count_matrix_weights(model)
        )

        result.update({
            "total_parameters": (
                total_parameters
            ),
            "matrix_weights": (
                matrix_weights
            ),
        })

        if bits_per_weight is not None:
            result[
                "theoretical_weight_size_mb"
            ] = theoretical_size_mb(
                matrix_weights,
                bits_per_weight,
            )

    return result


# ============================================================
# 7. 输出比较表
# ============================================================

def value_or_dash(
    value,
    number_format: str,
) -> str:
    if value is None:
        return "-"

    return format(value, number_format)


def print_comparison_table(
    results: list[dict[str, Any]],
) -> None:
    print("\n")
    print("=" * 140)
    print("不同算法场景性能比较")
    print("=" * 140)

    header = (
        f"{'模型':<18}"
        f"{'Accuracy':>11}"
        f"{'Precision':>11}"
        f"{'Recall':>11}"
        f"{'F1':>11}"
        f"{'总时间(s)':>13}"
        f"{'单张(ms)':>13}"
        f"{'吞吐(img/s)':>15}"
        f"{'文件(MB)':>12}"
        f"{'理论权重(MB)':>16}"
    )

    print(header)
    print("-" * 140)

    for result in results:
        total_time = result.get(
            "total_time",
            0.0,
        )

        time_per_image = result.get(
            "time_per_image",
            0.0,
        )

        throughput = (
            1.0 / time_per_image
            if time_per_image > 0
            else 0.0
        )

        theoretical_size = result.get(
            "theoretical_weight_size_mb"
        )

        print(
            f"{result['model']:<18}"
            f"{result['accuracy'] * 100:>10.2f}%"
            f"{result['precision']:>11.4f}"
            f"{result['recall']:>11.4f}"
            f"{result['f1']:>11.4f}"
            f"{total_time:>13.6f}"
            f"{time_per_image * 1000:>13.6f}"
            f"{throughput:>15.2f}"
            f"{result['model_file_size_mb']:>12.4f}"
            f"{value_or_dash(theoretical_size, '.4f'):>16}"
        )

    print("=" * 140)


def print_detailed_reports(
    results: list[dict[str, Any]],
    class_names: list[str],
) -> None:
    for result in results:
        print(
            f"\n{result['model']} 详细分类报告"
        )
        print("-" * 60)

        report = result[
            "classification_report"
        ]

        # classification_report已保存为字典，
        # 这里手动打印核心类别指标。
        print(
            f"{'类别':<20}"
            f"{'Precision':>12}"
            f"{'Recall':>12}"
            f"{'F1-score':>12}"
            f"{'Support':>10}"
        )

        for class_name in class_names:
            class_result = report[
                class_name
            ]

            print(
                f"{class_name:<20}"
                f"{class_result['precision']:>12.4f}"
                f"{class_result['recall']:>12.4f}"
                f"{class_result['f1-score']:>12.4f}"
                f"{int(class_result['support']):>10}"
            )




# ============================================================
# 8. 保存Excel
# ============================================================

EXCEL_DARK_BLUE = "1F4E78"
EXCEL_LIGHT_BLUE = "D9EAF7"
EXCEL_WHITE = "FFFFFF"
EXCEL_BORDER = "B7C9D6"


def _format_excel_header(
    worksheet,
    row_number: int = 1,
) -> None:
    """设置Excel表头格式。"""
    header_fill = PatternFill(
        fill_type="solid",
        fgColor=EXCEL_DARK_BLUE,
    )
    header_font = Font(
        color=EXCEL_WHITE,
        bold=True,
    )
    border = Border(
        left=Side(style="thin", color=EXCEL_BORDER),
        right=Side(style="thin", color=EXCEL_BORDER),
        top=Side(style="thin", color=EXCEL_BORDER),
        bottom=Side(style="thin", color=EXCEL_BORDER),
    )

    for cell in worksheet[row_number]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border

    worksheet.row_dimensions[row_number].height = 30


def _format_excel_body(
    worksheet,
    start_row: int,
    end_row: int,
) -> None:
    """使用浅蓝色和白色交替显示数据行。"""
    light_fill = PatternFill(
        fill_type="solid",
        fgColor=EXCEL_LIGHT_BLUE,
    )
    white_fill = PatternFill(
        fill_type="solid",
        fgColor=EXCEL_WHITE,
    )
    border = Border(
        left=Side(style="thin", color=EXCEL_BORDER),
        right=Side(style="thin", color=EXCEL_BORDER),
        top=Side(style="thin", color=EXCEL_BORDER),
        bottom=Side(style="thin", color=EXCEL_BORDER),
    )

    for row_number in range(start_row, end_row + 1):
        row_fill = (
            light_fill
            if row_number % 2 == 0
            else white_fill
        )

        for cell in worksheet[row_number]:
            cell.fill = row_fill
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )


def _autofit_excel_columns(
    worksheet,
    minimum_width: int = 10,
    maximum_width: int = 36,
) -> None:
    """根据单元格内容调整列宽，并限制最大宽度。"""
    for column_index in range(
        1,
        worksheet.max_column + 1,
    ):
        maximum_length = 0

        for row_index in range(
            1,
            worksheet.max_row + 1,
        ):
            value = worksheet.cell(
                row=row_index,
                column=column_index,
            ).value

            if value is None:
                continue

            maximum_length = max(
                maximum_length,
                len(str(value)),
            )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(
            max(maximum_length + 2, minimum_width),
            maximum_width,
        )


def _percent_for_excel(
    value,
) -> float | None:
    """
    将0~1比例或已经乘100的百分数，
    统一转换为Excel中显示的百分数数值。
    """
    if value is None:
        return None

    numeric_value = float(value)

    if 0.0 <= numeric_value <= 1.0:
        return numeric_value * 100.0

    return numeric_value


def save_excel_results(
    results: list[dict[str, Any]],
    output_dir: Path,
    dataset_root: Path,
    class_names: list[str],
    device: torch.device,
) -> Path:
    """
    将比较汇总和各类别分类报告保存到Excel。

    工作表：
    1. Summary：不同模型总体性能比较
    2. Class Details：每个模型、每个类别的详细指标
    3. Run Info：数据集与运行信息
    """
    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    excel_path = (
        output_dir
        / "comparison_results.xlsx"
    )

    workbook = Workbook()

    # --------------------------------------------------------
    # Sheet 1: Summary
    # --------------------------------------------------------

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    summary_headers = [
        "Model",
        "Operation",
        "Accuracy (%)",
        "Precision",
        "Recall",
        "F1-score",
        "Total Time (s)",
        "Time per Image (ms)",
        "Throughput (img/s)",
        "Total Parameters",
        "Matrix Weights",
        "Model File Size (MB)",
        "Theoretical Weight Size (MB)",
        "Best Epoch",
        "Best Validation Accuracy (%)",
        "Boolean Thresholds",
        "Patch Dimension",
        "Input Mode",
        "Boolean Channels",
        "Generated TM Input Shape",
        "Threshold Source",
    ]

    summary_sheet.append(
        summary_headers
    )

    for result in results:
        time_per_image = float(
            result.get(
                "time_per_image",
                0.0,
            )
            or 0.0
        )

        throughput = (
            1.0 / time_per_image
            if time_per_image > 0
            else 0.0
        )

        thresholds = (
            result.get("channel_thresholds")
            or result.get("rgb_channel_thresholds")
            or result.get("boolean_thresholds")
        )
        patch_dim = result.get(
            "patch_dim"
        )

        summary_sheet.append([
            result.get("model"),
            result.get("operation"),
            float(
                result.get(
                    "accuracy",
                    0.0,
                )
            ) * 100.0,
            result.get("precision"),
            result.get("recall"),
            result.get("f1"),
            result.get("total_time"),
            time_per_image * 1000.0,
            throughput,
            result.get("total_parameters"),
            result.get("matrix_weights"),
            result.get("model_file_size_mb"),
            result.get(
                "theoretical_weight_size_mb"
            ),
            result.get("best_epoch"),
            _percent_for_excel(
                result.get(
                    "best_validation_accuracy"
                )
            ),
            (
                str(thresholds)
                if thresholds is not None
                else None
            ),
            (
                str(patch_dim)
                if patch_dim is not None
                else None
            ),
            result.get("input_mode"),
            result.get("boolean_channels"),
            (
                str(
                    result.get(
                        "generated_tsetlin_input_shape"
                    )
                )
                if result.get(
                    "generated_tsetlin_input_shape"
                ) is not None
                else None
            ),
            result.get("threshold_source"),
        ])

    _format_excel_header(
        summary_sheet,
        row_number=1,
    )

    if summary_sheet.max_row >= 2:
        _format_excel_body(
            summary_sheet,
            start_row=2,
            end_row=summary_sheet.max_row,
        )

    summary_sheet.freeze_panes = "A2"
    summary_sheet.auto_filter.ref = (
        summary_sheet.dimensions
    )

    summary_sheet.column_dimensions["A"].width = 18
    summary_sheet.column_dimensions["B"].width = 32

    for cell in summary_sheet["C"][1:]:
        cell.number_format = "0.00"

    for column_letter in ("D", "E", "F"):
        for cell in summary_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0.0000"

    for column_letter in ("G", "H"):
        for cell in summary_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0.000000"

    for cell in summary_sheet["I"][1:]:
        cell.number_format = "0.00"

    for column_letter in ("J", "K", "N"):
        for cell in summary_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0"

    for column_letter in ("L", "M"):
        for cell in summary_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0.0000"

    for cell in summary_sheet["O"][1:]:
        cell.number_format = "0.00"

    _autofit_excel_columns(
        summary_sheet,
    )

    # --------------------------------------------------------
    # Sheet 2: Class Details
    # --------------------------------------------------------

    details_sheet = workbook.create_sheet(
        "Class Details"
    )

    details_headers = [
        "Model",
        "Class",
        "Precision",
        "Recall",
        "F1-score",
        "Support",
    ]

    details_sheet.append(
        details_headers
    )

    for result in results:
        report = result.get(
            "classification_report",
            {},
        )

        for class_name in class_names:
            class_result = report.get(
                class_name,
                {},
            )

            details_sheet.append([
                result.get("model"),
                class_name,
                class_result.get("precision"),
                class_result.get("recall"),
                class_result.get("f1-score"),
                class_result.get("support"),
            ])

        for average_name in (
            "macro avg",
            "weighted avg",
        ):
            average_result = report.get(
                average_name,
                {},
            )

            details_sheet.append([
                result.get("model"),
                average_name,
                average_result.get("precision"),
                average_result.get("recall"),
                average_result.get("f1-score"),
                average_result.get("support"),
            ])

    _format_excel_header(
        details_sheet,
        row_number=1,
    )

    if details_sheet.max_row >= 2:
        _format_excel_body(
            details_sheet,
            start_row=2,
            end_row=details_sheet.max_row,
        )

    details_sheet.freeze_panes = "A2"
    details_sheet.auto_filter.ref = (
        details_sheet.dimensions
    )

    for column_letter in ("C", "D", "E"):
        for cell in details_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0.0000"

    for cell in details_sheet["F"][1:]:
        cell.number_format = "0"

    _autofit_excel_columns(
        details_sheet,
    )

    # --------------------------------------------------------
    # Sheet 3: Run Info
    # --------------------------------------------------------

    info_sheet = workbook.create_sheet(
        "Run Info"
    )

    info_sheet.append([
        "Item",
        "Value",
    ])

    info_rows = [
        [
            "Generated Time",
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
        ],
        [
            "Dataset Root",
            str(dataset_root.resolve()),
        ],
        [
            "Image Size",
            str(IMAGE_SIZE),
        ],
        [
            "Class Count",
            len(class_names),
        ],
        [
            "Class Names",
            ", ".join(class_names),
        ],
        [
            "PyTorch Device",
            str(device),
        ],
        [
            "Tested Models",
            ", ".join(
                str(result.get("model"))
                for result in results
            ),
        ],
    ]

    for row in info_rows:
        info_sheet.append(row)

    _format_excel_header(
        info_sheet,
        row_number=1,
    )

    if info_sheet.max_row >= 2:
        _format_excel_body(
            info_sheet,
            start_row=2,
            end_row=info_sheet.max_row,
        )

    info_sheet.freeze_panes = "A2"
    info_sheet.column_dimensions["A"].width = 24
    info_sheet.column_dimensions["B"].width = 80
    info_sheet["B5"].alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )

    workbook.save(
        excel_path
    )

    return excel_path


# ============================================================
# 8. 保存CSV和JSON
# ============================================================

def save_results(
    results: list[dict[str, Any]],
    output_dir: Path,
    dataset_root: Path,
    class_names: list[str],
    device: torch.device,
) -> None:
    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    json_path = (
        output_dir
        / "comparison_results.json"
    )

    csv_path = (
        output_dir
        / "comparison_summary.csv"
    )

    json_data = {
        "dataset_root": str(
            dataset_root.resolve()
        ),
        "image_size": IMAGE_SIZE,
        "class_names": class_names,
        "torch_device": str(device),
        "results": results,
    }

    with json_path.open(
        "w",
        encoding="utf-8",
    ) as file:
        json.dump(
            json_data,
            file,
            indent=2,
            ensure_ascii=False,
        )

    csv_fields = [
        "model",
        "operation",
        "accuracy",
        "precision",
        "recall",
        "f1",
        "total_time",
        "time_per_image",
        "throughput_images_per_second",
        "total_parameters",
        "matrix_weights",
        "model_file_size_mb",
        "theoretical_weight_size_mb",
    ]

    with csv_path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=csv_fields,
        )

        writer.writeheader()

        for result in results:
            time_per_image = result.get(
                "time_per_image",
                0.0,
            )

            row = {
                field: result.get(field)
                for field in csv_fields
            }

            row[
                "throughput_images_per_second"
            ] = (
                1.0 / time_per_image
                if time_per_image > 0
                else 0.0
            )

            writer.writerow(row)

    excel_path = save_excel_results(
        results,
        output_dir,
        dataset_root,
        class_names,
        device,
    )

    print("\nComparison JSON:", json_path.resolve())
    print("Comparison CSV:", csv_path.resolve())
    print("Comparison Excel:", excel_path.resolve())


# ============================================================
# 9. 主程序
# ============================================================

def main() -> None:
    args = parse_arguments()
    set_seed()

    device = resolve_device(
        args.device
    )

    print("PyTorch device:", device)

    if (
        device.type == "cuda"
        and "tsetlin" in args.models
    ):
        print(
            "Warning: PyTorch模型使用CUDA，"
            "Tsetlin仍可能使用CPU，"
            "推理时间不能视为完全公平比较。"
        )

    # --------------------------------------------------------
    # MLP和BNN使用灰度测试数据
    # --------------------------------------------------------

    _, _, grayscale_testloader, class_names = (
        create_dataloaders(
            dataset_root=args.dataset_root,
            batch_size=args.batch_size,
            image_size=IMAGE_SIZE,
            input_mode="grayscale",
        )
    )

    grayscale_cached_batches, _ = (
        cache_test_batches(
            grayscale_testloader
        )
    )

    # --------------------------------------------------------
    # Tsetlin根据checkpoint自动选择输入模式
    # --------------------------------------------------------

    tsetlin_artifact = None
    tsetlin_cached_batches = None

    if "tsetlin" in args.models:

        tsetlin_artifact = load_tsetlin_artifact(
            args.tsetlin_model,
            class_names,
        )

        tsetlin_input_mode = (
            tsetlin_artifact["input_mode"]
        )

        print(
            "Tsetlin input mode from checkpoint:",
            tsetlin_input_mode,
        )

        (
            _,
            _,
            tsetlin_testloader,
            tsetlin_class_names,
        ) = create_dataloaders(
            dataset_root=args.dataset_root,
            batch_size=args.batch_size,
            image_size=IMAGE_SIZE,
            input_mode=tsetlin_input_mode,
        )

        if (
            list(tsetlin_class_names)
            != list(class_names)
        ):
            raise RuntimeError(
                "Tsetlin数据集和公共数据集的"
                "类别顺序不一致。"
            )

        tsetlin_cached_batches, _ = (
            cache_test_batches(
                tsetlin_testloader
            )
        )

    results: list[dict[str, Any]] = []

    # --------------------------------------------------------
    # MLP
    # --------------------------------------------------------

    if "mlp" in args.models:
        print("\nTesting MLP...")

        mlp_model = RGBDMLP(
            num_classes=len(class_names)
        ).to(device)

        load_torch_weights(
            mlp_model,
            args.mlp_model,
            device,
        )

        mlp_metrics = evaluate_torch_model(
            mlp_model,
            grayscale_cached_batches,
            class_names,
            device,
        )

        results.append(
            add_common_statistics(
                model_name="MLP",
                model_path=args.mlp_model,
                metrics=mlp_metrics,
                model=mlp_model,
                bits_per_weight=32,
                operation="FP32 Linear",
            )
        )

    # --------------------------------------------------------
    # BNN：BinaryLinear
    # --------------------------------------------------------

    if "bnn" in args.models:
        print("\nTesting BNN...")

        bnn_model = RGBDBNN(
            num_classes=len(class_names),
            use_xnor_popcount=False,
        ).to(device)

        load_torch_weights(
            bnn_model,
            args.bnn_model,
            device,
        )

        bnn_metrics = evaluate_torch_model(
            bnn_model,
            grayscale_cached_batches,
            class_names,
            device,
        )

        results.append(
            add_common_statistics(
                model_name="BNN",
                model_path=args.bnn_model,
                metrics=bnn_metrics,
                model=bnn_model,
                bits_per_weight=1,
                operation=(
                    "Binary weight dot product"
                ),
            )
        )

    # --------------------------------------------------------
    # BNN：XNOR + Popcount
    # --------------------------------------------------------

    if "bnn_xnor" in args.models:
        print(
            "\nTesting BNN XNOR+Popcount..."
        )

        bnn_xnor_model = RGBDBNN(
            num_classes=len(class_names),
            use_xnor_popcount=True,
        ).to(device)

        load_torch_weights(
            bnn_xnor_model,
            args.bnn_model,
            device,
        )

        bnn_xnor_metrics = (
            evaluate_torch_model(
                bnn_xnor_model,
                grayscale_cached_batches,
                class_names,
                device,
            )
        )

        results.append(
            add_common_statistics(
                model_name="BNN-XNOR",
                model_path=args.bnn_model,
                metrics=bnn_xnor_metrics,
                model=bnn_xnor_model,
                bits_per_weight=1,
                operation="XNOR + Popcount",
            )
        )

    # --------------------------------------------------------
    # Tsetlin Machine
    # --------------------------------------------------------

    if "tsetlin" in args.models:
        print("\nTesting Tsetlin Machine...")

        tsetlin_metrics, metadata = (
            evaluate_tsetlin_model(
                tsetlin_artifact,
                tsetlin_cached_batches,
                class_names,
            )
        )

        tsetlin_result = (
            add_common_statistics(
                model_name="Tsetlin",
                model_path=(
                    args.tsetlin_model
                ),
                metrics=tsetlin_metrics,
                operation=(
                    "Clause-based Boolean inference"
                ),
            )
        )

        tsetlin_result.update(
            metadata
        )

        results.append(
            tsetlin_result
        )

    if not results:
        raise RuntimeError(
            "没有选择任何可测试模型。"
        )

    print_comparison_table(results)

    print_detailed_reports(
        results,
        class_names,
    )

    save_results(
        results,
        args.output_dir,
        args.dataset_root,
        class_names,
        device,
    )


if __name__ == "__main__":
    main()
