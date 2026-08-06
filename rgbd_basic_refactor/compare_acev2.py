# -*- coding: utf-8 -*-
"""
公共模型比较文件。

默认比较四种算法场景：
1. MLP
2. BNN（BinaryLinear 二值点积）
3. BNN（XNOR + Popcount 推理）
4. Tsetlin Machine

所有模型使用同一个 RGBDObjectDataset 测试集。
默认使用 CPU，便于与 CPU Tsetlin Machine 进行较公平的比较。

运行：
    python compare.py

只比较指定模型：
    python compare.py --models mlp bnn
    python compare.py --models mlp bnn bnn_xnor tsetlin

选择Tsetlin输入模式：
    python compare.py --tsetlin-input-mode grayscale
    python compare.py --tsetlin-input-mode rgb

默认使用auto：优先读取Tsetlin checkpoint中的input_mode；
旧checkpoint没有input_mode时，根据阈值字段推断，无法推断则使用grayscale。

ACEv2输出两种口径：
1. Current ACEv2：按当前Python推理实现的数据类型估算。
2. Theoretical ACEv2：按目标模型的理论位宽估算。

位宽口径：
- MLP：FP32激活 x FP32权重。
- BNN：当前BinaryLinear仍执行FP32线性运算；理论值使用
  FP32激活 x 1-bit权重。
- BNN-XNOR：第一层使用FP32输入 x 1-bit权重，后续层使用
  1-bit激活 x 1-bit权重。
- Tsetlin：使用1-bit clause逻辑操作上界估算。ACEv2论文没有
  原生定义clause AND/NOT，因此该值标记为兼容性估算。

指定模型路径：
    python compare.py ^
        --mlp-model models/rgbd_mlp_2.pth ^
        --bnn-model models/rgbd_bnn_2.pth ^
        --tsetlin-model tsetlin_output/best_tsetlin_rgbd.pkl
"""

from __future__ import annotations

import argparse
import csv
import json
import pickle
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import torch
import torch.nn as nn
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.metrics import classification_report

from bnn import RGBDBNN
from common import (
    BATCH_SIZE,
    DATASET_ROOT,
    IMAGE_SIZE,
    build_classification_metrics,
    create_dataloaders,
    set_seed,
)
from mlp import RGBDMLP


SCRIPT_DIR = Path(__file__).resolve().parent

DEFAULT_MLP_MODEL = (
    SCRIPT_DIR
    / "models"
    / "rgbd_mlp_2.pth"
)

DEFAULT_BNN_MODEL = (
    SCRIPT_DIR
    / "models"
    / "rgbd_bnn_2.pth"
)

DEFAULT_TSETLIN_MODEL = (
    SCRIPT_DIR
    / "tsetlin_output"
    / "best_tsetlin_rgbd.pkl"
)

DEFAULT_OUTPUT_DIR = (
    SCRIPT_DIR
    / "comparison_results"
)

SUPPORTED_MODELS = (
    "mlp",
    "bnn",
    "bnn_xnor",
    "tsetlin",
)


# ============================================================
# 1. 命令行参数
# ============================================================

def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Compare MLP, BNN, BNN-XNOR/Popcount "
            "and Tsetlin Machine on the same test set."
        )
    )

    parser.add_argument(
        "--models",
        nargs="+",
        choices=SUPPORTED_MODELS,
        default=list(SUPPORTED_MODELS),
        help=(
            "要比较的模型。默认："
            "mlp bnn bnn_xnor tsetlin"
        ),
    )

    parser.add_argument(
        "--dataset-root",
        type=Path,
        default=DATASET_ROOT,
        help="RGB-D Object Dataset 根目录。",
    )

    parser.add_argument(
        "--batch-size",
        type=int,
        default=BATCH_SIZE,
        help="测试 batch size。",
    )

    parser.add_argument(
        "--device",
        choices=("cpu", "cuda", "auto"),
        default="cpu",
        help=(
            "PyTorch模型使用的设备。默认CPU，"
            "便于与CPU Tsetlin公平比较。"
        ),
    )

    parser.add_argument(
        "--mlp-model",
        type=Path,
        default=DEFAULT_MLP_MODEL,
        help="MLP .pth 文件。",
    )

    parser.add_argument(
        "--bnn-model",
        type=Path,
        default=DEFAULT_BNN_MODEL,
        help="BNN .pth 文件。",
    )

    parser.add_argument(
        "--tsetlin-model",
        type=Path,
        default=DEFAULT_TSETLIN_MODEL,
        help="Tsetlin .pkl 文件。",
    )

    parser.add_argument(
        "--tsetlin-input-mode",
        choices=("auto", "grayscale", "rgb"),
        default="auto",
        help=(
            "Tsetlin输入模式。auto优先读取checkpoint；"
            "也可以手动选择grayscale或rgb。"
        ),
    )

    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="比较结果保存目录。",
    )

    return parser.parse_args()


def resolve_device(device_name: str) -> torch.device:
    if device_name == "cpu":
        return torch.device("cpu")

    if device_name == "cuda":
        if not torch.cuda.is_available():
            raise RuntimeError(
                "指定了CUDA，但当前环境没有可用CUDA设备。"
            )
        return torch.device("cuda")

    return torch.device(
        "cuda"
        if torch.cuda.is_available()
        else "cpu"
    )


# ============================================================
# 2. 模型参数加载
# ============================================================

def load_torch_weights(
    model: nn.Module,
    model_path: Path,
    device: torch.device,
) -> None:
    """
    兼容三种常见保存形式：

    1. torch.save(model.state_dict(), path)
    2. {"model_state_dict": state_dict}
    3. {"state_dict": state_dict}
    """
    if not model_path.exists():
        raise FileNotFoundError(
            f"模型文件不存在：{model_path.resolve()}"
        )

    checkpoint = torch.load(
        model_path,
        map_location=device,
        weights_only=False,
    )

    if isinstance(checkpoint, nn.Module):
        state_dict = checkpoint.state_dict()

    elif isinstance(checkpoint, dict):
        if "model_state_dict" in checkpoint:
            state_dict = checkpoint[
                "model_state_dict"
            ]
        elif "state_dict" in checkpoint:
            state_dict = checkpoint["state_dict"]
        else:
            state_dict = checkpoint

    else:
        raise ValueError(
            f"无法识别模型文件格式：{model_path}"
        )

    model.load_state_dict(state_dict)
    model.eval()


# ============================================================
# 3. 测试数据缓存
# ============================================================

def cache_test_batches(
    testloader,
) -> tuple[
    list[tuple[torch.Tensor, torch.Tensor]],
    np.ndarray,
]:
    """
    先把测试集读取到内存。

    这样计时时不包含硬盘读取和PIL预处理时间，
    更接近模型自身的推理时间。
    """
    cached_batches = []
    labels_all = []

    for images, labels in testloader:
        cached_batches.append((
            images.contiguous(),
            labels.contiguous(),
        ))
        labels_all.append(
            labels.numpy()
        )

    if not cached_batches:
        raise RuntimeError(
            "测试集中没有读取到任何数据。"
        )

    return (
        cached_batches,
        np.concatenate(
            labels_all,
            axis=0,
        ),
    )


# ============================================================
# 4. PyTorch模型测试
# ============================================================

def evaluate_torch_model(
    model: nn.Module,
    cached_batches,
    class_names: list[str],
    device: torch.device,
) -> dict[str, Any]:
    predictions_all = []
    labels_all = []

    model.eval()

    # 预热一次，减少首次调用初始化对时间的影响。
    with torch.no_grad():
        first_images, _ = cached_batches[0]
        _ = model(
            first_images.to(device)
        )

    if device.type == "cuda":
        torch.cuda.synchronize()

    start_time = time.perf_counter()

    with torch.no_grad():
        for images, labels in cached_batches:
            images = images.to(
                device,
                non_blocking=True,
            )

            outputs = model(images)

            predictions = torch.argmax(
                outputs,
                dim=1,
            )

            predictions_all.extend(
                predictions.cpu().numpy()
            )
            labels_all.extend(
                labels.numpy()
            )

    if device.type == "cuda":
        torch.cuda.synchronize()

    total_time = (
        time.perf_counter()
        - start_time
    )

    return build_classification_metrics(
        labels_all,
        predictions_all,
        class_names,
        total_time=total_time,
    )


# ============================================================
# 5. Tsetlin测试数据转换
# ============================================================

def _normalise_tsetlin_input_mode(
    value: Any,
) -> str | None:
    if value is None:
        return None

    normalised = str(value).strip().lower()

    aliases = {
        "gray": "grayscale",
        "grey": "grayscale",
        "greyscale": "grayscale",
        "grayscale": "grayscale",
        "rgb": "rgb",
    }

    return aliases.get(normalised)


def resolve_tsetlin_input_mode(
    model_path: Path,
    requested_mode: str,
) -> str:
    """
    决定Tsetlin使用灰度还是RGB输入。

    auto模式优先读取checkpoint中的input_mode。对于没有该字段的
    旧checkpoint，如果存在rgb_channel_thresholds则判断为RGB；
    其他情况默认使用当前项目正式比较所用的grayscale。
    """
    if requested_mode != "auto":
        return requested_mode

    if not model_path.exists():
        raise FileNotFoundError(
            "Tsetlin模型文件不存在："
            f"{model_path.resolve()}"
        )

    with model_path.open("rb") as file:
        checkpoint = pickle.load(file)

    if isinstance(checkpoint, dict):
        saved_mode = _normalise_tsetlin_input_mode(
            checkpoint.get("input_mode")
        )

        if saved_mode is not None:
            return saved_mode

        if "rgb_channel_thresholds" in checkpoint:
            return "rgb"

    return "grayscale"


def get_tsetlin_threshold_config(
    checkpoint: Any,
    input_mode: str,
) -> tuple[tuple[float, ...], str, str]:
    """
    返回：阈值、阈值排列方式、阈值来源字段。

    grayscale：一个灰度通道分别通过全部阈值。
    rgb/per_channel：R、G、B各自使用一个阈值。
    rgb/channel_then_threshold：每个RGB通道分别通过全部阈值。
    """
    if not isinstance(checkpoint, dict):
        return (-0.5, 0.0, 0.5), (
            "all_thresholds"
            if input_mode == "grayscale"
            else "per_channel"
        ), "default"

    if input_mode == "grayscale":
        candidate_keys = (
            "boolean_thresholds",
            "grayscale_thresholds",
            "channel_thresholds",
            "thresholds",
        )
    else:
        candidate_keys = (
            "rgb_channel_thresholds",
            "boolean_thresholds",
            "channel_thresholds",
            "thresholds",
        )

    threshold_values = None
    threshold_source = "default"

    for key in candidate_keys:
        values = checkpoint.get(key)
        if values is not None:
            threshold_values = values
            threshold_source = key
            break

    if threshold_values is None:
        threshold_values = (-0.5, 0.0, 0.5)

    thresholds = tuple(
        float(value)
        for value in threshold_values
    )

    if not thresholds:
        raise ValueError(
            "Tsetlin checkpoint中的Boolean阈值为空。"
        )

    if input_mode == "grayscale":
        threshold_layout = "all_thresholds"
    elif (
        threshold_source in (
            "rgb_channel_thresholds",
            "default",
        )
        and len(thresholds) == 3
    ):
        threshold_layout = "per_channel"
    else:
        threshold_layout = "channel_then_threshold"

    return (
        thresholds,
        threshold_layout,
        threshold_source,
    )


def cached_batches_to_boolean_numpy(
    cached_batches,
    input_mode: str,
    thresholds: tuple[float, ...],
    threshold_layout: str,
) -> tuple[np.ndarray, np.ndarray]:

    feature_batches = []
    label_batches = []

    with torch.no_grad():
        for images, labels in cached_batches:

            expected_channels = (
                1
                if input_mode == "grayscale"
                else 3
            )

            if (
                images.ndim != 4
                or images.shape[1] != expected_channels
            ):
                raise ValueError(
                    f"Tsetlin {input_mode}模型需要"
                    f"[B,{expected_channels},H,W]输入，"
                    f"当前为{tuple(images.shape)}"
                )

            if input_mode == "grayscale":
                boolean_planes = [
                    (
                        images[:, 0, :, :] > threshold
                    ).to(torch.uint32)
                    for threshold in thresholds
                ]

                boolean_tensor = torch.stack(
                    boolean_planes,
                    dim=-1,
                ).contiguous()

            elif threshold_layout == "per_channel":
                if len(thresholds) != 3:
                    raise ValueError(
                        "RGB per_channel模式需要3个阈值，"
                        f"当前为{len(thresholds)}个。"
                    )

                threshold_tensor = torch.tensor(
                    thresholds,
                    dtype=images.dtype,
                    device=images.device,
                ).view(1, 3, 1, 1)

                boolean_tensor = (
                    images > threshold_tensor
                ).to(torch.uint32).permute(
                    0,
                    2,
                    3,
                    1,
                ).contiguous()

            else:
                boolean_planes = []

                # 排列顺序：R的全部阈值、G的全部阈值、B的全部阈值。
                for channel_index in range(3):
                    for threshold in thresholds:
                        boolean_planes.append(
                            (
                                images[
                                    :,
                                    channel_index,
                                    :,
                                    :,
                                ] > threshold
                            ).to(torch.uint32)
                        )

                boolean_tensor = torch.stack(
                    boolean_planes,
                    dim=-1,
                ).contiguous()

            feature_batches.append(
                boolean_tensor.cpu().numpy().astype(
                    np.uint32,
                    copy=False,
                )
            )

            label_batches.append(
                labels.cpu().numpy().astype(
                    np.uint32,
                    copy=False,
                )
            )

    X_test = np.concatenate(
        feature_batches,
        axis=0,
    )

    y_test = np.concatenate(
        label_batches,
        axis=0,
    )

    return (
        np.ascontiguousarray(
            X_test,
            dtype=np.uint32,
        ),
        np.ascontiguousarray(
            y_test,
            dtype=np.uint32,
        ),
    )


def evaluate_tsetlin_model(
    model_path: Path,
    cached_batches,
    class_names: list[str],
    input_mode: str,
) -> tuple[dict[str, Any], dict[str, Any]]:
    if not model_path.exists():
        raise FileNotFoundError(
            "Tsetlin模型文件不存在："
            f"{model_path.resolve()}"
        )

    # pickle只能加载可信的模型文件。
    with model_path.open("rb") as file:
        checkpoint = pickle.load(file)

    if isinstance(checkpoint, dict):
        if "model" not in checkpoint:
            raise KeyError(
                "Tsetlin checkpoint中没有'model'。"
            )

        model = checkpoint["model"]

        saved_input_mode = (
            _normalise_tsetlin_input_mode(
                checkpoint.get("input_mode")
            )
        )

        if (
            saved_input_mode is not None
            and saved_input_mode != input_mode
        ):
            print(
                "Warning: 手动选择的Tsetlin输入模式"
                f"({input_mode})与checkpoint记录"
                f"({saved_input_mode})不同。"
            )

        saved_classes = checkpoint.get(
            "class_names"
        )

        if (
            saved_classes is not None
            and list(saved_classes)
            != list(class_names)
        ):
            raise RuntimeError(
                "Tsetlin模型类别顺序与测试集不一致。"
            )

        (
            thresholds,
            threshold_layout,
            threshold_source,
        ) = get_tsetlin_threshold_config(
            checkpoint,
            input_mode,
        )

        model_metadata = {
            "input_mode": input_mode,
            "best_epoch": checkpoint.get(
                "epoch"
            ),
            "best_validation_accuracy": (
                checkpoint.get(
                    "validation_accuracy"
                )
            ),
            "boolean_thresholds": list(
                thresholds
            ),
            "threshold_layout": threshold_layout,
            "threshold_source": threshold_source,
            "patch_dim": checkpoint.get(
                "patch_dim"
            ),
        }

    else:
        # 兼容直接pickle.dump(model)
        model = checkpoint
        (
            thresholds,
            threshold_layout,
            threshold_source,
        ) = get_tsetlin_threshold_config(
            checkpoint,
            input_mode,
        )
        model_metadata = {
            "input_mode": input_mode,
            "boolean_thresholds": list(
                thresholds
            ),
            "threshold_layout": threshold_layout,
            "threshold_source": threshold_source,
        }

    X_test, y_test = (
        cached_batches_to_boolean_numpy(
            cached_batches,
            input_mode,
            thresholds,
            threshold_layout,
        )
    )

    if isinstance(checkpoint, dict):
        saved_boolean_channels = checkpoint.get(
            "boolean_channels"
        )

        if (
            saved_boolean_channels is not None
            and X_test.shape[-1]
            != int(saved_boolean_channels)
        ):
            raise RuntimeError(
                "Boolean通道数与checkpoint不一致："
                f"测试输入={X_test.shape[-1]}，"
                f"checkpoint={saved_boolean_channels}。"
            )

    # 预热一次。
    warmup_count = min(
        1,
        len(X_test),
    )

    if warmup_count > 0:
        _ = model.predict(
            X_test[:warmup_count]
        )

    start_time = time.perf_counter()

    predictions = model.predict(X_test)

    total_time = (
        time.perf_counter()
        - start_time
    )

    metrics = build_classification_metrics(
        y_test,
        predictions,
        class_names,
        total_time=total_time,
    )

    model_metadata.update(
        calculate_tsetlin_acev2(
            model=model,
            checkpoint=checkpoint,
            X_test=X_test,
            class_count=len(class_names),
        )
    )

    return metrics, model_metadata


# ============================================================
# 6. 参数、权重和文件大小
# ============================================================

def count_total_parameters(
    model: nn.Module,
) -> int:
    return sum(
        parameter.numel()
        for parameter in model.parameters()
    )


def count_matrix_weights(
    model: nn.Module,
) -> int:
    """
    只统计Linear等二维权重，
    不统计bias和BatchNorm的一维参数。
    """
    return sum(
        parameter.numel()
        for name, parameter
        in model.named_parameters()
        if (
            "weight" in name
            and parameter.dim() >= 2
        )
    )


def get_file_size_mb(
    file_path: Path,
) -> float:
    if not file_path.exists():
        return 0.0

    return (
        file_path.stat().st_size
        / (1024 * 1024)
    )


def theoretical_size_mb(
    weight_count: int,
    bits_per_weight: int,
) -> float:
    return (
        weight_count
        * bits_per_weight
        / 8
        / (1024 * 1024)
    )


def calculate_neural_acev2(
    model: nn.Module,
    activation_bits: int,
    weight_bits: int,
    first_layer_activation_bits: int | None = None,
    include_bias: bool = True,
    include_batchnorm: bool = True,
    float_bias: bool = True,
    batchnorm_bits: int = 32,
) -> dict[str, float | int]:
    """
    计算全连接MLP/BNN单张图片推理的ACEv2。

    计算范围：
    1. Linear层MAC：
       MAC数量 * activation_bits * weight_bits
    2. Linear层bias加法。
    3. BatchNorm推理：y = scale * x + offset。

    ReLU比较、argmax、数据预处理和内存访问不在
    ACEv2的标准计算范围内，因此这里不统计。
    """
    total_macs = 0
    total_bias_adds = 0
    total_batchnorm_elements = 0
    mac_acev2 = 0

    linear_index = 0

    for module in model.modules():
        if isinstance(module, nn.Linear):
            layer_macs = (
                module.in_features
                * module.out_features
            )

            layer_activation_bits = activation_bits

            if (
                linear_index == 0
                and first_layer_activation_bits
                is not None
            ):
                layer_activation_bits = (
                    first_layer_activation_bits
                )

            total_macs += layer_macs
            mac_acev2 += (
                layer_macs
                * layer_activation_bits
                * weight_bits
            )

            if include_bias and module.bias is not None:
                total_bias_adds += module.out_features

            linear_index += 1

        elif (
            include_batchnorm
            and isinstance(module, nn.BatchNorm1d)
        ):
            total_batchnorm_elements += (
                module.num_features
            )

    if float_bias:
        bias_cost_per_add = 6 * 32
    else:
        bias_cost_per_add = max(
            activation_bits,
            weight_bits,
        )

    bias_acev2 = total_bias_adds * bias_cost_per_add

    # FP32 BatchNorm：每个元素一次乘法和一次浮点加法。
    batchnorm_mul_cost = (
        batchnorm_bits * batchnorm_bits
        - batchnorm_bits
    )
    batchnorm_add_cost = 6 * batchnorm_bits
    batchnorm_acev2 = (
        total_batchnorm_elements
        * (
            batchnorm_mul_cost
            + batchnorm_add_cost
        )
    )

    total_acev2 = (
        mac_acev2
        + bias_acev2
        + batchnorm_acev2
    )

    return {
        "macs_per_image": total_macs,
        "bias_adds_per_image": total_bias_adds,
        "batchnorm_elements_per_image": (
            total_batchnorm_elements
        ),
        "mac_acev2": mac_acev2,
        "bias_acev2": bias_acev2,
        "batchnorm_acev2": batchnorm_acev2,
        "acev2": total_acev2,
        "acev2_million": total_acev2 / 1e6,
        "acev2_billion": total_acev2 / 1e9,
    }


def attach_acev2_profiles(
    result: dict[str, Any],
    current_profile: dict[str, float | int],
    theoretical_profile: dict[str, float | int],
    basis: str,
) -> None:
    """把当前实现和理论位宽两种ACEv2写入结果。"""
    result.update(current_profile)

    for key, value in theoretical_profile.items():
        if key in (
            "macs_per_image",
            "bias_adds_per_image",
            "batchnorm_elements_per_image",
        ):
            continue

        result[f"theoretical_{key}"] = value

    result["acev2_basis"] = basis


def _first_positive_int(*values: Any) -> int | None:
    for value in values:
        if value is None:
            continue

        try:
            number = int(value)
        except (TypeError, ValueError):
            continue

        if number > 0:
            return number

    return None


def calculate_tsetlin_acev2(
    model: Any,
    checkpoint: Any,
    X_test: np.ndarray,
    class_count: int,
) -> dict[str, Any]:
    """
    生成Tsetlin的ACEv2兼容性估算。

    ACEv2论文没有为Tsetlin clause逻辑定义标准成本。本函数把
    每次1-bit AND/NOT或1-bit vote加法记为1个ACEv2单位，并按
    所有literal均被检查的最坏情况计算clause evaluation上界。
    因此该值适合在同一项目内作理论参考，不能解释为真实能耗。
    """
    checkpoint_dict = (
        checkpoint
        if isinstance(checkpoint, dict)
        else {}
    )

    config = checkpoint_dict.get("config", {})
    if not isinstance(config, dict):
        config = {}

    number_of_clauses = _first_positive_int(
        checkpoint_dict.get("number_of_clauses"),
        checkpoint_dict.get("clauses"),
        config.get("number_of_clauses"),
        config.get("clauses"),
        getattr(model, "number_of_clauses", None),
    )

    number_of_features = _first_positive_int(
        checkpoint_dict.get("number_of_features"),
        config.get("number_of_features"),
        getattr(model, "number_of_features", None),
        int(np.prod(X_test.shape[1:])),
    )

    number_of_patches = _first_positive_int(
        checkpoint_dict.get("number_of_patches"),
        config.get("number_of_patches"),
        getattr(model, "number_of_patches", None),
        1,
    )

    resolved_class_count = _first_positive_int(
        checkpoint_dict.get("number_of_classes"),
        config.get("number_of_classes"),
        getattr(model, "number_of_classes", None),
        class_count,
    )

    if (
        number_of_clauses is None
        or number_of_features is None
        or number_of_patches is None
        or resolved_class_count is None
    ):
        return {
            "acev2": None,
            "acev2_million": None,
            "acev2_billion": None,
            "theoretical_acev2": None,
            "theoretical_acev2_million": None,
            "theoretical_acev2_billion": None,
            "acev2_basis": (
                "Tsetlin estimate unavailable: checkpoint "
                "does not expose clause/feature counts"
            ),
        }

    literal_count = 2 * number_of_features
    clause_evaluations = (
        resolved_class_count
        * number_of_clauses
        * number_of_patches
    )

    clause_logic_operations = (
        clause_evaluations
        * max(literal_count - 1, 0)
    )

    # 每个class对正/负clause vote进行1-bit理论累加。
    vote_additions = (
        resolved_class_count
        * number_of_clauses
    )

    total_acev2 = (
        clause_logic_operations
        + vote_additions
    )

    return {
        "logic_operations_per_image": (
            clause_logic_operations
        ),
        "clause_evaluations_per_image": (
            clause_evaluations
        ),
        "tsetlin_number_of_features": (
            number_of_features
        ),
        "tsetlin_number_of_clauses": (
            number_of_clauses
        ),
        "tsetlin_number_of_patches": (
            number_of_patches
        ),
        "acev2": total_acev2,
        "acev2_million": total_acev2 / 1e6,
        "acev2_billion": total_acev2 / 1e9,
        "theoretical_acev2": total_acev2,
        "theoretical_acev2_million": (
            total_acev2 / 1e6
        ),
        "theoretical_acev2_billion": (
            total_acev2 / 1e9
        ),
        "acev2_basis": (
            "ACEv2-compatible upper bound: 1-bit clause "
            "logic plus 1-bit vote additions"
        ),
    }


def add_common_statistics(
    model_name: str,
    model_path: Path,
    metrics: dict[str, Any],
    model: nn.Module | None = None,
    bits_per_weight: int | None = None,
    operation: str = "",
) -> dict[str, Any]:
    result = {
        "model": model_name,
        "operation": operation,
        **metrics,
        "model_file_size_mb": (
            get_file_size_mb(
                model_path
            )
        ),
    }

    if model is not None:
        total_parameters = (
            count_total_parameters(model)
        )
        matrix_weights = (
            count_matrix_weights(model)
        )

        result.update({
            "total_parameters": (
                total_parameters
            ),
            "matrix_weights": (
                matrix_weights
            ),
        })

        if bits_per_weight is not None:
            result[
                "theoretical_weight_size_mb"
            ] = theoretical_size_mb(
                matrix_weights,
                bits_per_weight,
            )

    return result


# ============================================================
# 7. 输出比较表
# ============================================================

def value_or_dash(
    value,
    number_format: str,
) -> str:
    if value is None:
        return "-"

    return format(value, number_format)


def print_comparison_table(
    results: list[dict[str, Any]],
) -> None:
    print("\n")
    print("=" * 177)
    print("不同算法场景性能比较")
    print("=" * 177)

    header = (
        f"{'模型':<18}"
        f"{'Accuracy':>11}"
        f"{'Precision':>11}"
        f"{'Recall':>11}"
        f"{'F1':>11}"
        f"{'总时间(s)':>13}"
        f"{'单张(ms)':>13}"
        f"{'吞吐(img/s)':>15}"
        f"{'文件(MB)':>12}"
        f"{'理论权重(MB)':>16}"
        f"{'当前ACEv2(M)':>16}"
        f"{'理论ACEv2(M)':>16}"
    )

    print(header)
    print("-" * 177)

    for result in results:
        total_time = result.get(
            "total_time",
            0.0,
        )

        time_per_image = result.get(
            "time_per_image",
            0.0,
        )

        throughput = (
            1.0 / time_per_image
            if time_per_image > 0
            else 0.0
        )

        theoretical_size = result.get(
            "theoretical_weight_size_mb"
        )

        acev2_million = result.get(
            "acev2_million"
        )

        theoretical_acev2_million = result.get(
            "theoretical_acev2_million"
        )

        print(
            f"{result['model']:<18}"
            f"{result['accuracy'] * 100:>10.2f}%"
            f"{result['precision']:>11.4f}"
            f"{result['recall']:>11.4f}"
            f"{result['f1']:>11.4f}"
            f"{total_time:>13.6f}"
            f"{time_per_image * 1000:>13.6f}"
            f"{throughput:>15.2f}"
            f"{result['model_file_size_mb']:>12.4f}"
            f"{value_or_dash(theoretical_size, '.4f'):>16}"
            f"{value_or_dash(acev2_million, '.6f'):>16}"
            f"{value_or_dash(theoretical_acev2_million, '.6f'):>16}"
        )

    print("=" * 177)


def print_detailed_reports(
    results: list[dict[str, Any]],
    class_names: list[str],
) -> None:
    for result in results:
        print(
            f"\n{result['model']} 详细分类报告"
        )
        print("-" * 60)

        report = result[
            "classification_report"
        ]

        # classification_report已保存为字典，
        # 这里手动打印核心类别指标。
        print(
            f"{'类别':<20}"
            f"{'Precision':>12}"
            f"{'Recall':>12}"
            f"{'F1-score':>12}"
            f"{'Support':>10}"
        )

        for class_name in class_names:
            class_result = report[
                class_name
            ]

            print(
                f"{class_name:<20}"
                f"{class_result['precision']:>12.4f}"
                f"{class_result['recall']:>12.4f}"
                f"{class_result['f1-score']:>12.4f}"
                f"{int(class_result['support']):>10}"
            )




# ============================================================
# 8. 保存Excel
# ============================================================

EXCEL_DARK_BLUE = "1F4E78"
EXCEL_LIGHT_BLUE = "D9EAF7"
EXCEL_WHITE = "FFFFFF"
EXCEL_BORDER = "B7C9D6"


def _format_excel_header(
    worksheet,
    row_number: int = 1,
) -> None:
    """设置Excel表头格式。"""
    header_fill = PatternFill(
        fill_type="solid",
        fgColor=EXCEL_DARK_BLUE,
    )
    header_font = Font(
        color=EXCEL_WHITE,
        bold=True,
    )
    border = Border(
        left=Side(style="thin", color=EXCEL_BORDER),
        right=Side(style="thin", color=EXCEL_BORDER),
        top=Side(style="thin", color=EXCEL_BORDER),
        bottom=Side(style="thin", color=EXCEL_BORDER),
    )

    for cell in worksheet[row_number]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border

    worksheet.row_dimensions[row_number].height = 30


def _format_excel_body(
    worksheet,
    start_row: int,
    end_row: int,
) -> None:
    """使用浅蓝色和白色交替显示数据行。"""
    light_fill = PatternFill(
        fill_type="solid",
        fgColor=EXCEL_LIGHT_BLUE,
    )
    white_fill = PatternFill(
        fill_type="solid",
        fgColor=EXCEL_WHITE,
    )
    border = Border(
        left=Side(style="thin", color=EXCEL_BORDER),
        right=Side(style="thin", color=EXCEL_BORDER),
        top=Side(style="thin", color=EXCEL_BORDER),
        bottom=Side(style="thin", color=EXCEL_BORDER),
    )

    for row_number in range(start_row, end_row + 1):
        row_fill = (
            light_fill
            if row_number % 2 == 0
            else white_fill
        )

        for cell in worksheet[row_number]:
            cell.fill = row_fill
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )


def _autofit_excel_columns(
    worksheet,
    minimum_width: int = 10,
    maximum_width: int = 36,
) -> None:
    """根据单元格内容调整列宽，并限制最大宽度。"""
    for column_index in range(
        1,
        worksheet.max_column + 1,
    ):
        maximum_length = 0

        for row_index in range(
            1,
            worksheet.max_row + 1,
        ):
            value = worksheet.cell(
                row=row_index,
                column=column_index,
            ).value

            if value is None:
                continue

            maximum_length = max(
                maximum_length,
                len(str(value)),
            )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(
            max(maximum_length + 2, minimum_width),
            maximum_width,
        )


def _percent_for_excel(
    value,
) -> float | None:
    """
    将0~1比例或已经乘100的百分数，
    统一转换为Excel中显示的百分数数值。
    """
    if value is None:
        return None

    numeric_value = float(value)

    if 0.0 <= numeric_value <= 1.0:
        return numeric_value * 100.0

    return numeric_value


def save_excel_results(
    results: list[dict[str, Any]],
    output_dir: Path,
    dataset_root: Path,
    class_names: list[str],
    device: torch.device,
) -> Path:
    """
    将比较汇总和各类别分类报告保存到Excel。

    工作表：
    1. Summary：不同模型总体性能比较
    2. Class Details：每个模型、每个类别的详细指标
    3. Run Info：数据集与运行信息
    """
    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    excel_path = (
        output_dir
        / "comparison_results.xlsx"
    )

    workbook = Workbook()

    # --------------------------------------------------------
    # Sheet 1: Summary
    # --------------------------------------------------------

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    summary_headers = [
        "Model",
        "Operation",
        "Accuracy (%)",
        "Precision",
        "Recall",
        "F1-score",
        "Total Time (s)",
        "Time per Image (ms)",
        "Throughput (img/s)",
        "Total Parameters",
        "Matrix Weights",
        "MACs per Image",
        "Bias Adds per Image",
        "BatchNorm Elements per Image",
        "Current ACEv2 (Million)",
        "Theoretical ACEv2 (Million)",
        "Current ACEv2 (Billion)",
        "Theoretical ACEv2 (Billion)",
        "ACEv2 Basis",
        "Model File Size (MB)",
        "Theoretical Weight Size (MB)",
        "Best Epoch",
        "Best Validation Accuracy (%)",
        "Boolean Thresholds",
        "Patch Dimension",
        "Input Mode",
        "Threshold Layout",
        "Threshold Source",
    ]

    summary_sheet.append(
        summary_headers
    )

    for result in results:
        time_per_image = float(
            result.get(
                "time_per_image",
                0.0,
            )
            or 0.0
        )

        throughput = (
            1.0 / time_per_image
            if time_per_image > 0
            else 0.0
        )

        thresholds = result.get(
            "boolean_thresholds"
        )
        patch_dim = result.get(
            "patch_dim"
        )

        summary_sheet.append([
            result.get("model"),
            result.get("operation"),
            float(
                result.get(
                    "accuracy",
                    0.0,
                )
            ) * 100.0,
            result.get("precision"),
            result.get("recall"),
            result.get("f1"),
            result.get("total_time"),
            time_per_image * 1000.0,
            throughput,
            result.get("total_parameters"),
            result.get("matrix_weights"),
            result.get("macs_per_image"),
            result.get("bias_adds_per_image"),
            result.get(
                "batchnorm_elements_per_image"
            ),
            result.get("acev2_million"),
            result.get(
                "theoretical_acev2_million"
            ),
            result.get("acev2_billion"),
            result.get(
                "theoretical_acev2_billion"
            ),
            result.get("acev2_basis"),
            result.get("model_file_size_mb"),
            result.get(
                "theoretical_weight_size_mb"
            ),
            result.get("best_epoch"),
            _percent_for_excel(
                result.get(
                    "best_validation_accuracy"
                )
            ),
            (
                str(thresholds)
                if thresholds is not None
                else None
            ),
            (
                str(patch_dim)
                if patch_dim is not None
                else None
            ),
            result.get("input_mode"),
            result.get("threshold_layout"),
            result.get("threshold_source"),
        ])

    _format_excel_header(
        summary_sheet,
        row_number=1,
    )

    if summary_sheet.max_row >= 2:
        _format_excel_body(
            summary_sheet,
            start_row=2,
            end_row=summary_sheet.max_row,
        )

    summary_sheet.freeze_panes = "A2"
    summary_sheet.auto_filter.ref = (
        summary_sheet.dimensions
    )

    summary_sheet.column_dimensions["A"].width = 18
    summary_sheet.column_dimensions["B"].width = 32

    for cell in summary_sheet["C"][1:]:
        cell.number_format = "0.00"

    for column_letter in ("D", "E", "F"):
        for cell in summary_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0.0000"

    for column_letter in ("G", "H"):
        for cell in summary_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0.000000"

    for cell in summary_sheet["I"][1:]:
        cell.number_format = "0.00"

    for column_letter in ("J", "K", "L", "M", "N", "V"):
        for cell in summary_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0"

    for column_letter in ("O", "P", "Q", "R"):
        for cell in summary_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0.000000"

    for column_letter in ("T", "U"):
        for cell in summary_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0.0000"

    for cell in summary_sheet["W"][1:]:
        cell.number_format = "0.00"

    _autofit_excel_columns(
        summary_sheet,
    )

    # --------------------------------------------------------
    # Sheet 2: Class Details
    # --------------------------------------------------------

    details_sheet = workbook.create_sheet(
        "Class Details"
    )

    details_headers = [
        "Model",
        "Class",
        "Precision",
        "Recall",
        "F1-score",
        "Support",
    ]

    details_sheet.append(
        details_headers
    )

    for result in results:
        report = result.get(
            "classification_report",
            {},
        )

        for class_name in class_names:
            class_result = report.get(
                class_name,
                {},
            )

            details_sheet.append([
                result.get("model"),
                class_name,
                class_result.get("precision"),
                class_result.get("recall"),
                class_result.get("f1-score"),
                class_result.get("support"),
            ])

        for average_name in (
            "macro avg",
            "weighted avg",
        ):
            average_result = report.get(
                average_name,
                {},
            )

            details_sheet.append([
                result.get("model"),
                average_name,
                average_result.get("precision"),
                average_result.get("recall"),
                average_result.get("f1-score"),
                average_result.get("support"),
            ])

    _format_excel_header(
        details_sheet,
        row_number=1,
    )

    if details_sheet.max_row >= 2:
        _format_excel_body(
            details_sheet,
            start_row=2,
            end_row=details_sheet.max_row,
        )

    details_sheet.freeze_panes = "A2"
    details_sheet.auto_filter.ref = (
        details_sheet.dimensions
    )

    for column_letter in ("C", "D", "E"):
        for cell in details_sheet[
            column_letter
        ][1:]:
            cell.number_format = "0.0000"

    for cell in details_sheet["F"][1:]:
        cell.number_format = "0"

    _autofit_excel_columns(
        details_sheet,
    )

    # --------------------------------------------------------
    # Sheet 3: Run Info
    # --------------------------------------------------------

    info_sheet = workbook.create_sheet(
        "Run Info"
    )

    info_sheet.append([
        "Item",
        "Value",
    ])

    info_rows = [
        [
            "Generated Time",
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            ),
        ],
        [
            "Dataset Root",
            str(dataset_root.resolve()),
        ],
        [
            "Image Size",
            str(IMAGE_SIZE),
        ],
        [
            "Class Count",
            len(class_names),
        ],
        [
            "Class Names",
            ", ".join(class_names),
        ],
        [
            "PyTorch Device",
            str(device),
        ],
        [
            "Tested Models",
            ", ".join(
                str(result.get("model"))
                for result in results
            ),
        ],
    ]

    for row in info_rows:
        info_sheet.append(row)

    _format_excel_header(
        info_sheet,
        row_number=1,
    )

    if info_sheet.max_row >= 2:
        _format_excel_body(
            info_sheet,
            start_row=2,
            end_row=info_sheet.max_row,
        )

    info_sheet.freeze_panes = "A2"
    info_sheet.column_dimensions["A"].width = 24
    info_sheet.column_dimensions["B"].width = 80
    info_sheet["B5"].alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )

    workbook.save(
        excel_path
    )

    return excel_path


# ============================================================
# 8. 保存CSV和JSON
# ============================================================

def save_results(
    results: list[dict[str, Any]],
    output_dir: Path,
    dataset_root: Path,
    class_names: list[str],
    device: torch.device,
) -> None:
    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    json_path = (
        output_dir
        / "comparison_results.json"
    )

    csv_path = (
        output_dir
        / "comparison_summary.csv"
    )

    json_data = {
        "dataset_root": str(
            dataset_root.resolve()
        ),
        "image_size": IMAGE_SIZE,
        "class_names": class_names,
        "torch_device": str(device),
        "results": results,
    }

    with json_path.open(
        "w",
        encoding="utf-8",
    ) as file:
        json.dump(
            json_data,
            file,
            indent=2,
            ensure_ascii=False,
        )

    csv_fields = [
        "model",
        "operation",
        "accuracy",
        "precision",
        "recall",
        "f1",
        "total_time",
        "time_per_image",
        "throughput_images_per_second",
        "total_parameters",
        "matrix_weights",
        "macs_per_image",
        "bias_adds_per_image",
        "batchnorm_elements_per_image",
        "logic_operations_per_image",
        "clause_evaluations_per_image",
        "mac_acev2",
        "bias_acev2",
        "batchnorm_acev2",
        "acev2",
        "acev2_million",
        "acev2_billion",
        "theoretical_mac_acev2",
        "theoretical_bias_acev2",
        "theoretical_batchnorm_acev2",
        "theoretical_acev2",
        "theoretical_acev2_million",
        "theoretical_acev2_billion",
        "acev2_basis",
        "model_file_size_mb",
        "theoretical_weight_size_mb",
        "tsetlin_number_of_features",
        "tsetlin_number_of_clauses",
        "tsetlin_number_of_patches",
        "input_mode",
        "boolean_thresholds",
        "threshold_layout",
        "threshold_source",
    ]

    with csv_path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=csv_fields,
        )

        writer.writeheader()

        for result in results:
            time_per_image = result.get(
                "time_per_image",
                0.0,
            )

            row = {
                field: result.get(field)
                for field in csv_fields
            }

            row[
                "throughput_images_per_second"
            ] = (
                1.0 / time_per_image
                if time_per_image > 0
                else 0.0
            )

            writer.writerow(row)

    excel_path = save_excel_results(
        results,
        output_dir,
        dataset_root,
        class_names,
        device,
    )

    print("\nComparison JSON:", json_path.resolve())
    print("Comparison CSV:", csv_path.resolve())
    print("Comparison Excel:", excel_path.resolve())


# ============================================================
# 9. 主程序
# ============================================================

def main() -> None:
    args = parse_arguments()
    set_seed()

    device = resolve_device(
        args.device
    )

    print("PyTorch device:", device)

    if (
        device.type == "cuda"
        and "tsetlin" in args.models
    ):
        print(
            "Warning: PyTorch模型使用CUDA，"
            "Tsetlin仍可能使用CPU，"
            "推理时间不能视为完全公平比较。"
        )

    # --------------------------------------------------------
    # MLP和BNN使用灰度测试数据
    # --------------------------------------------------------

    _, _, grayscale_testloader, class_names = (
        create_dataloaders(
            dataset_root=args.dataset_root,
            batch_size=args.batch_size,
            image_size=IMAGE_SIZE,
            input_mode="grayscale",
        )
    )

    grayscale_cached_batches, _ = (
        cache_test_batches(
            grayscale_testloader
        )
    )

    # --------------------------------------------------------
    # Tsetlin根据命令行或checkpoint选择灰度/RGB数据
    # --------------------------------------------------------

    tsetlin_cached_batches = None
    tsetlin_input_mode = None

    if "tsetlin" in args.models:
        tsetlin_input_mode = (
            resolve_tsetlin_input_mode(
                args.tsetlin_model,
                args.tsetlin_input_mode,
            )
        )

        print(
            "Tsetlin input mode:",
            tsetlin_input_mode,
        )

        if tsetlin_input_mode == "grayscale":
            # 与MLP/BNN复用完全相同的灰度测试batch。
            tsetlin_cached_batches = (
                grayscale_cached_batches
            )

        else:
            _, _, rgb_testloader, rgb_class_names = (
                create_dataloaders(
                    dataset_root=args.dataset_root,
                    batch_size=args.batch_size,
                    image_size=IMAGE_SIZE,
                    input_mode="rgb",
                )
            )

            if list(rgb_class_names) != list(class_names):
                raise RuntimeError(
                    "RGB和灰度数据集的类别顺序不一致。"
                )

            tsetlin_cached_batches, _ = (
                cache_test_batches(
                    rgb_testloader
                )
            )

    results: list[dict[str, Any]] = []

    # --------------------------------------------------------
    # MLP
    # --------------------------------------------------------

    if "mlp" in args.models:
        print("\nTesting MLP...")

        mlp_model = RGBDMLP(
            num_classes=len(class_names)
        ).to(device)

        load_torch_weights(
            mlp_model,
            args.mlp_model,
            device,
        )

        mlp_metrics = evaluate_torch_model(
            mlp_model,
            grayscale_cached_batches,
            class_names,
            device,
        )

        mlp_result = add_common_statistics(
            model_name="MLP",
            model_path=args.mlp_model,
            metrics=mlp_metrics,
            model=mlp_model,
            bits_per_weight=32,
            operation="FP32 Linear",
        )

        mlp_current_acev2 = calculate_neural_acev2(
            mlp_model,
            activation_bits=32,
            weight_bits=32,
            include_bias=True,
        )

        # MLP的当前实现与理论模型均为FP32。
        mlp_theoretical_acev2 = (
            calculate_neural_acev2(
                mlp_model,
                activation_bits=32,
                weight_bits=32,
                include_bias=True,
            )
        )

        attach_acev2_profiles(
            mlp_result,
            current_profile=mlp_current_acev2,
            theoretical_profile=(
                mlp_theoretical_acev2
            ),
            basis=(
                "Current and theoretical: FP32 activation "
                "x FP32 weight; FP32 bias"
            ),
        )

        results.append(mlp_result)

        print(
            "MLP MACs per image:",
            mlp_current_acev2["macs_per_image"],
        )
        print(
            "MLP bias adds per image:",
            mlp_current_acev2[
                "bias_adds_per_image"
            ],
        )
        print(
            "MLP current/theoretical ACEv2:",
            f'{mlp_current_acev2["acev2_million"]:.6f}',
            "/",
            f'{mlp_theoretical_acev2["acev2_million"]:.6f}',
            "Million",
        )

    # --------------------------------------------------------
    # BNN：BinaryLinear
    # --------------------------------------------------------

    if "bnn" in args.models:
        print("\nTesting BNN...")

        bnn_model = RGBDBNN(
            num_classes=len(class_names),
            use_xnor_popcount=False,
        ).to(device)

        load_torch_weights(
            bnn_model,
            args.bnn_model,
            device,
        )

        bnn_metrics = evaluate_torch_model(
            bnn_model,
            grayscale_cached_batches,
            class_names,
            device,
        )

        bnn_result = add_common_statistics(
            model_name="BNN",
            model_path=args.bnn_model,
            metrics=bnn_metrics,
            model=bnn_model,
            bits_per_weight=1,
            operation=(
                "Binary weight dot product"
            ),
        )

        # 当前BinaryLinear仍通过PyTorch FP32 Linear执行。
        bnn_current_acev2 = calculate_neural_acev2(
            bnn_model,
            activation_bits=32,
            weight_bits=32,
            include_bias=True,
            include_batchnorm=True,
        )

        # 理论权重为1 bit；当前模型激活仍为FP32。
        bnn_theoretical_acev2 = (
            calculate_neural_acev2(
                bnn_model,
                activation_bits=32,
                weight_bits=1,
                include_bias=True,
                include_batchnorm=True,
            )
        )

        attach_acev2_profiles(
            bnn_result,
            current_profile=bnn_current_acev2,
            theoretical_profile=(
                bnn_theoretical_acev2
            ),
            basis=(
                "Current: FP32 PyTorch Linear with binary-valued "
                "weights; theoretical: FP32 activation x 1-bit "
                "weight; FP32 BatchNorm/bias"
            ),
        )

        results.append(bnn_result)

        print(
            "BNN current/theoretical ACEv2:",
            f'{bnn_current_acev2["acev2_million"]:.6f}',
            "/",
            f'{bnn_theoretical_acev2["acev2_million"]:.6f}',
            "Million",
        )

    # --------------------------------------------------------
    # BNN：XNOR + Popcount
    # --------------------------------------------------------

    if "bnn_xnor" in args.models:
        print(
            "\nTesting BNN XNOR+Popcount..."
        )

        bnn_xnor_model = RGBDBNN(
            num_classes=len(class_names),
            use_xnor_popcount=True,
        ).to(device)

        load_torch_weights(
            bnn_xnor_model,
            args.bnn_model,
            device,
        )

        bnn_xnor_metrics = (
            evaluate_torch_model(
                bnn_xnor_model,
                grayscale_cached_batches,
                class_names,
                device,
            )
        )

        bnn_xnor_result = add_common_statistics(
            model_name="BNN-XNOR",
            model_path=args.bnn_model,
            metrics=bnn_xnor_metrics,
            model=bnn_xnor_model,
            bits_per_weight=1,
            operation="XNOR + Popcount",
        )

        # 原始灰度像素是FP32，第一层不能直接使用1-bit激活；
        # 第一层后才进入XNOR/Popcount的1-bit激活路径。
        bnn_xnor_current_acev2 = (
            calculate_neural_acev2(
                bnn_xnor_model,
                activation_bits=1,
                weight_bits=1,
                first_layer_activation_bits=32,
                include_bias=True,
                include_batchnorm=True,
            )
        )

        # 全二值硬件的理想理论值：所有层均为1 x 1 bit。
        bnn_xnor_theoretical_acev2 = (
            calculate_neural_acev2(
                bnn_xnor_model,
                activation_bits=1,
                weight_bits=1,
                include_bias=True,
                include_batchnorm=True,
            )
        )

        attach_acev2_profiles(
            bnn_xnor_result,
            current_profile=(
                bnn_xnor_current_acev2
            ),
            theoretical_profile=(
                bnn_xnor_theoretical_acev2
            ),
            basis=(
                "Current algorithm: first layer FP32 activation "
                "x 1-bit weight, later 1 x 1 bit; theoretical: "
                "all Linear layers 1 x 1 bit; FP32 BatchNorm/bias"
            ),
        )

        results.append(bnn_xnor_result)

        print(
            "BNN-XNOR current/theoretical ACEv2:",
            f'{bnn_xnor_current_acev2["acev2_million"]:.6f}',
            "/",
            f'{bnn_xnor_theoretical_acev2["acev2_million"]:.6f}',
            "Million",
        )

    # --------------------------------------------------------
    # Tsetlin Machine
    # --------------------------------------------------------

    if "tsetlin" in args.models:
        print("\nTesting Tsetlin Machine...")

        tsetlin_metrics, metadata = (
            evaluate_tsetlin_model(
                args.tsetlin_model,
                tsetlin_cached_batches,
                class_names,
                tsetlin_input_mode,
            )
        )

        tsetlin_result = (
            add_common_statistics(
                model_name="Tsetlin",
                model_path=(
                    args.tsetlin_model
                ),
                metrics=tsetlin_metrics,
                operation=(
                    "Clause-based Boolean inference"
                ),
            )
        )

        tsetlin_result.update(
            metadata
        )

        print(
            "Tsetlin ACEv2-compatible estimate:",
            value_or_dash(
                tsetlin_result.get(
                    "acev2_million"
                ),
                ".6f",
            ),
            "Million",
        )

        results.append(
            tsetlin_result
        )

    if not results:
        raise RuntimeError(
            "没有选择任何可测试模型。"
        )

    print_comparison_table(results)

    print_detailed_reports(
        results,
        class_names,
    )

    save_results(
        results,
        args.output_dir,
        args.dataset_root,
        class_names,
        device,
    )


if __name__ == "__main__":
    main()
